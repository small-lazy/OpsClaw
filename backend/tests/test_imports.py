import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from opsweaver import store
from opsweaver.api import app


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(store, 'DATA', tmp_path)
    monkeypatch.setattr(store, 'DB', tmp_path / 'opsweaver.sqlite3')
    with TestClient(app) as client:
        client.post('/api/v1/workspace/session')
        yield client


def test_batch_analysis_encoding_errors_and_persistence(client):
    response = client.post('/api/v1/imports', files=[('files', ('中文.csv', '城市,销售额\n成都,12\n成都,12\n上海,\n'.encode('gb18030'), 'text/csv')), ('files', ('broken.json', b'{broken', 'application/json')), ('files', ('events.jsonl', b'{"id":1,"payload":{"a":2}}\n{"id":2}', 'application/json'))])
    assert response.status_code == 200
    batch = response.json()['data']
    assert batch['status'] == 'partial'
    assert batch['dataset_count'] == 2
    assert batch['files'][1]['status'] == 'failed'
    dataset_id = batch['files'][0]['dataset_ids'][0]
    data = client.get('/api/v1/datasets/' + dataset_id).json()['data']
    assert data['analysis']['duplicate_rows'] == 1
    assert data['analysis']['missing_cells'] == 1
    assert data['columns'][1]['numeric']['sum'] == 24
    assert client.get('/api/v1/sources/' + data['source_id'] + '/rows?offset=1&limit=1').json()['data']['rows'] == [{'城市': '成都', '销售额': '12'}]
    store.initialize()
    assert client.get('/api/v1/imports/' + batch['id']).json()['data'] == batch
    assert client.get('/api/v1/datasets/' + dataset_id + '/rows?offset=-1').status_code == 422


def test_xlsx_multisheet_and_formula_text(client):
    book = Workbook()
    sheet = book.active
    sheet.title = '销售'
    sheet.append(['地区', '金额'])
    sheet.append(['华东', 35])
    book.create_sheet('空白')
    other = book.create_sheet('备注')
    other.append(['说明'])
    other.append(['=1+1'])
    content = io.BytesIO()
    book.save(content)
    batch = client.post('/api/v1/imports', files={'files': ('book.xlsx', content.getvalue())}).json()['data']
    assert batch['status'] == 'completed'
    assert batch['dataset_count'] == 2
    dataset = client.get('/api/v1/datasets/' + batch['files'][0]['dataset_ids'][1]).json()['data']
    assert dataset['preview'][0]['说明'] == '=1+1'
    assert any('公式' in text for text in dataset['analysis']['insights'])


def test_malformed_and_unsupported_files_are_independent(client):
    batch = client.post('/api/v1/imports', files=[('files', ('bad.xlsx', b'bad')), ('files', ('bad.csv', b'a,a\n1,2')), ('files', ('bad.exe', b'bad')), ('files', ('good.tsv', b'id\tvalue\n1\t4'))]).json()['data']
    assert batch['status'] == 'partial'
    assert [f['status'] for f in batch['files']] == ['failed', 'failed', 'failed', 'completed']
    client.cookies.clear()
    assert client.post('/api/v1/imports', files={'files': ('test.csv', b'a\n1')}).status_code == 401


def test_empty_standard_table_can_declare_zero_records(client):
    schema = client.get('/api/v1/business-schema').json()['data']['tables']
    columns = next(t['columns'] for t in schema if t['role'] == 'contacts')
    batch = client.post('/api/v1/imports', files={'files': ('contacts.csv', (','.join(columns) + '\n').encode())}).json()['data']
    assert batch['status'] == 'completed'
    dataset = client.get('/api/v1/datasets/' + batch['files'][0]['dataset_ids'][0]).json()['data']
    assert dataset['row_count'] == 0
    assert dataset['business']['role'] == 'contacts'


def test_csv_customer_unknown_consent_and_empty_roles_activate(client):
    schema = client.get('/api/v1/business-schema').json()['data']['tables']
    columns = next(t['columns'] for t in schema if t['role'] == 'customers')
    text = ','.join(columns) + '\nUSER-X,2026-01-01,regular,,0,2026-09-01,1,,2026-09-01\n'
    batch = client.post('/api/v1/imports', files={'files': ('customers.csv', text.encode())}).json()['data']
    dataset_id = batch['files'][0]['dataset_ids'][0]
    activated = client.post('/api/v1/datasets/' + dataset_id + '/activate')
    assert activated.status_code == 200, activated.text
    assert client.get('/api/v1/customers/USER-X/business').json()['data']['marketing_consent'] is None
    for table in schema:
        if table['role'] == 'customers':
            continue
        batch = client.post('/api/v1/imports', files={'files': (table['role'] + '.csv', (','.join(table['columns']) + '\n').encode())}).json()['data']
        activated = client.post('/api/v1/datasets/' + batch['files'][0]['dataset_ids'][0] + '/activate')
        assert activated.status_code == 200, activated.text
    assert client.get('/api/v1/customers/USER-X/business').json()['data']['data_fresh'] is True
