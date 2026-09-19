from __future__ import annotations

import csv
import io
import json
import math
import secrets
import sqlite3
import zipfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from xml.etree.ElementTree import ParseError
from openpyxl.utils.exceptions import InvalidFileException

from . import business, store
from .service import DomainError, require

MAX_FILE = 20 * 1024 * 1024
MAX_ROWS = 100000
MAX_COLUMNS = 200


def decode(raw):
    for encoding in ('utf-8-sig', 'utf-16' if raw.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8', 'gb18030'):
        try:
            return raw.decode(encoding)
        except UnicodeError:
            continue
    raise ValueError('无法识别文件编码，请保存为 UTF-8 或 GB18030。')


def cell(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False, allow_nan=False)
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError('数据包含非有限数值。')
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def table(headers, records):
    columns = [str(h).strip() if h is not None else '' for h in headers]
    if not columns or len(columns) > MAX_COLUMNS or any(not c for c in columns) or len(set(columns)) != len(columns):
        raise ValueError('表头不能为空或重复，每表最多 200 列。')
    rows = []
    for values in records:
        if not any(v is not None and v != '' for v in values):
            continue
        if len(values) != len(columns):
            raise ValueError(f'第 {len(rows) + 2} 行列数与表头不一致。')
        rows.append(dict(zip(columns, (cell(v) for v in values))))
        if len(rows) > MAX_ROWS:
            raise ValueError('每张表最多 100,000 行。')
    if not rows and not detect_role(columns):
        raise ValueError('文件没有数据行。')
    return columns, rows


def parse(name, raw):
    suffix = Path(name).suffix.lower()
    if suffix in ('.csv', '.tsv'):
        text = decode(raw)
        csv.field_size_limit(2 * 1024 * 1024)
        delimiter = '\t' if suffix == '.tsv' else ','
        if suffix == '.csv':
            try:
                delimiter = csv.Sniffer().sniff(text[:16384], delimiters=',;\t').delimiter
            except csv.Error:
                pass
        reader = csv.reader(io.StringIO(text, newline=''), delimiter=delimiter, strict=True)
        headers = next(reader, [])
        yield None, *table(headers, reader)
    elif suffix in ('.json', '.jsonl', '.ndjson'):
        text = decode(raw)
        records = [json.loads(line) for line in text.splitlines() if line.strip()] if suffix != '.json' else json.loads(text)
        if isinstance(records, dict) and isinstance(records.get('data'), list):
            records = records['data']
        if not isinstance(records, list) or not records or not all(isinstance(r, dict) for r in records):
            raise ValueError('JSON 需要对象数组（或 data 对象数组）；JSONL 每行一个对象。')
        if len(records) > MAX_ROWS:
            raise ValueError('每张表最多 100,000 行。')
        columns = list(dict.fromkeys(k for row in records for k in row))
        yield None, *table(columns, ([r.get(k) for k in columns] for r in records))
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Excel 解压后超过 100 MB。')
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
        try:
            if len(workbook.worksheets) > 30:
                raise ValueError('每个 Excel 最多 30 个工作表。')
            found = False
            for sheet in workbook.worksheets:
                if sheet.max_column and sheet.max_column > MAX_COLUMNS:
                    raise ValueError('每张表最多 200 列。')
                records = sheet.iter_rows(values_only=True)
                headers = next(records, None)
                if not headers or not any(v is not None for v in headers):
                    continue
                try:
                    parsed = table(headers, records)
                except ValueError as error:
                    if str(error) == '文件没有数据行。':
                        continue
                    raise
                found = True
                yield sheet.title, *parsed
            if not found:
                raise ValueError('工作簿没有可导入的数据表。')
        finally:
            workbook.close()
    else:
        raise ValueError('支持 CSV、TSV、XLSX、JSON、JSONL 文件。')


def profile(columns, rows):
    result, insights = [], []
    missing_cells = 0
    for name in columns:
        values = [r[name] for r in rows if r[name] is not None and str(r[name]).strip() != '']
        missing = len(rows) - len(values)
        missing_cells += missing
        counts = Counter(str(v) for v in values)
        nums = []
        for value in values:
            try:
                n = float(value)
                if isinstance(value, bool) or not math.isfinite(n):
                    break
                nums.append(n)
            except (TypeError, ValueError, OverflowError):
                break
        numeric = bool(values) and len(nums) == len(values)
        kind = 'number' if numeric else 'boolean' if values and all(isinstance(v, bool) for v in values) else 'text'
        col = {'name': name, 'type': kind, 'missing': missing, 'missing_percent': round(missing / max(len(rows), 1) * 100, 2), 'unique': len(counts), 'top_values': [{'value': v[:200], 'count': n} for v, n in counts.most_common(5)]}
        if numeric:
            # Avoid overflow when summing individually valid, large numbers.
            total = sum(nums)
            col['numeric'] = {'min': min(nums), 'max': max(nums), 'mean': sum(v / len(nums) for v in nums), 'sum': total if math.isfinite(total) else None}
            col['numeric'] = {k: v if v is None or math.isfinite(v) else None for k, v in col['numeric'].items()}
        if missing:
            insights.append(f'{name} 缺失 {missing} 个值（{col["missing_percent"]}%）。')
        result.append(col)
    duplicates = len(rows) - len({json.dumps(r, sort_keys=True, ensure_ascii=False) for r in rows})
    if duplicates:
        insights.append(f'发现 {duplicates} 行完全重复记录，统计仍保留这些记录。')
    if any(isinstance(v, str) and v.startswith('=') for r in rows for v in r.values()):
        insights.append('公式按文本保留，不执行公式；需要计算结果时请先在表格软件中转为值。')
    if not insights:
        insights.append('未发现缺失单元格或完全重复记录。')
    return result, {'summary': f'已分析 {len(rows):,} 行、{len(columns)} 列，其中 {sum(c["type"] == "number" for c in result)} 列为数值。', 'duplicate_rows': duplicates, 'missing_cells': missing_cells, 'insights': insights[:30]}


def detect_role(columns):
    with business._connect(store.DATA) as db:
        for role in business.ROLES:
            required = {r[1] for r in db.execute(f'PRAGMA table_info({role})')}
            if required and required.issubset(columns):
                return role
    return None


def save_dataset(connection, name, raw, sheet, columns, rows):
    dataset_id = 'dataset-' + secrets.token_hex(12)
    source_id = 'upload-' + secrets.token_hex(12)
    stats, analysis = profile(columns, rows)
    role = detect_role(columns)
    dataset = {'id': dataset_id, 'name': name + (f' · {sheet}' if sheet else ''), 'source_id': source_id, 'format': Path(name).suffix.lower()[1:], 'sheet': sheet, 'row_count': len(rows), 'columns': stats, 'preview': rows[:20], 'analysis': analysis, 'business': {'role': role, 'status': 'available' if role else 'generic', 'message': '标准字段完整，可接入业务分析。' if role else '已完成通用分析；接入业务规则需提供完整标准字段。'}, 'created_at': store.now()}
    store.put(connection, 'datasets', dataset)
    connection.executemany('INSERT INTO dataset_rows(dataset_id,row_index,document) VALUES(?,?,?)', ((dataset_id, i, json.dumps(row, ensure_ascii=False)) for i, row in enumerate(rows)))
    source = {'id': source_id, 'role': role or 'general', 'name': dataset['name'], 'rows': len(rows), 'columns': columns, 'preview': [{k: '' if v is None else str(v)[:1000] for k, v in r.items()} for r in rows[:20]], 'status': 'ready', 'coverage': False, 'as_of': dataset['created_at'], 'origin': 'user_upload', 'analysis_scope': 'dataset', 'dataset_id': dataset_id}
    store.put(connection, 'sources', source)
    return dataset


def import_files(files):
    batch = {'id': 'import-' + secrets.token_hex(12), 'created_at': store.now(), 'status': 'completed', 'files': [], 'dataset_count': 0}
    for name, raw in files:
        entry = {'name': name, 'status': 'completed', 'dataset_ids': []}
        try:
            if len(raw) > MAX_FILE:
                raise ValueError('单个文件不能超过 20 MB。')
            parsed = list(parse(name, raw))
            with store.transaction() as connection:
                for sheet, columns, rows in parsed:
                    dataset = save_dataset(connection, name, raw, sheet, columns, rows)
                    entry['dataset_ids'].append(dataset['id'])
                store.audit(connection, 'import.file', f'已导入 {name}，生成 {len(parsed)} 个数据集。')
            batch['dataset_count'] += len(parsed)
        except (ValueError, UnicodeError, csv.Error, zipfile.BadZipFile, InvalidFileException, ParseError, OSError, KeyError, TypeError, OverflowError) as error:
            entry.update(status='failed', dataset_ids=[], error=str(error)[:400])
        batch['files'].append(entry)
    failures = sum(f['status'] == 'failed' for f in batch['files'])
    batch['status'] = 'failed' if failures == len(files) else 'partial' if failures else 'completed'
    with store.transaction() as connection:
        store.put(connection, 'imports', batch)
    return batch


def rows(connection, dataset_id, offset, limit):
    if offset < 0 or not 1 <= limit <= 100:
        raise DomainError('INVALID_PAGE', 'offset 不能小于 0，limit 需为 1–100。', 422)
    dataset = require(connection, 'datasets', dataset_id)
    values = [json.loads(r[0]) for r in connection.execute('SELECT document FROM dataset_rows WHERE dataset_id=? ORDER BY row_index LIMIT ? OFFSET ?', (dataset_id, limit, offset))]
    return {'rows': values, 'total': dataset['row_count'], 'offset': offset, 'limit': limit}


def activate(dataset_id):
    with store.transaction() as connection:
        dataset = require(connection, 'datasets', dataset_id)
        role = dataset['business']['role']
        if role not in business.ROLES:
            raise DomainError('BUSINESS_FIELDS_REQUIRED', '请使用完整的标准业务表字段后再接入业务分析。', 422)
        first = not store.meta(connection, 'external_business')
        if first and role != 'customers':
            raise DomainError('CUSTOMERS_REQUIRED', '请先接入客户表，再接入关联业务表。', 422)
        records = [json.loads(r[0]) for r in connection.execute('SELECT document FROM dataset_rows WHERE dataset_id=? ORDER BY row_index', (dataset_id,))]
        path = store.DATA / 'user_business.sqlite3'
        business.create_empty(path, store.now())
        # An attached database makes business records, activation metadata and dataset state one transaction.
        connection.execute('ATTACH DATABASE ? AS imported', (str(path),))
        info = list(connection.execute(f'PRAGMA imported.table_info({role})'))
        fields = [r[1] for r in info]
        primary = fields[0]
        ids = set()
        for record in records:
            if any(record.get(k) is None or str(record[k]).strip() == '' for k in fields if k not in ('marketing_consent', 'coverage_reason', 'external_ref')):
                raise DomainError('INVALID_BUSINESS_ROW', '标准业务字段存在空值，请补全后重新导入。', 422)
            if record[primary] in ids:
                raise DomainError('DUPLICATE_BUSINESS_ID', '业务表主键重复，请先去重。', 422)
            ids.add(record[primary])
            for field in fields:
                if field.endswith('_at') or field == 'event_time':
                    try:
                        record[field] = business._time(str(record[field])).isoformat()
                    except (ValueError, TypeError) as error:
                        raise DomainError('INVALID_BUSINESS_DATE', f'{field} 需要 ISO 日期时间。', 422) from error
                if field.endswith('_minor') or field in ('marketing_consent', 'do_not_contact', 'coverage_complete'):
                    if field == 'marketing_consent' and (record[field] is None or str(record[field]).strip() == ''):
                        record[field] = None
                        continue
                    try:
                        value = int(record[field])
                        if float(record[field]) != value or value < 0 or value > 2**63 - 1 or (not field.endswith('_minor') and value not in (0, 1)):
                            raise ValueError()
                        record[field] = value
                    except (ValueError, TypeError, OverflowError) as error:
                        raise DomainError('INVALID_BUSINESS_NUMBER', f'{field} 必须为有效的非负整数，布尔字段使用 0 或 1。', 422) from error
            if role == 'orders' and (record['currency'] != 'CNY' or record['refunded_amount_minor'] > record['paid_amount_minor']):
                raise DomainError('INVALID_ORDER_AMOUNT', '订单币种需为 CNY，退款金额不能大于支付金额。', 422)
            if role != 'customers' and not connection.execute('SELECT 1 FROM imported.customers WHERE customer_id=?', (record['customer_id'],)).fetchone():
                raise DomainError('CUSTOMER_REFERENCE_MISSING', f'找不到关联客户 {record["customer_id"]}，请先接入对应客户。', 422)
        active_roles = set(store.meta(connection, 'external_roles') or []) | {role}
        placeholders = ','.join('?' for _ in fields)
        updates = ','.join(f'{k}=excluded.{k}' for k in fields[1:])
        connection.executemany(f'INSERT INTO imported.{role}({",".join(fields)}) VALUES({placeholders}) ON CONFLICT({primary}) DO UPDATE SET {updates}', ([r[k] for k in fields] for r in records))
        connection.execute("INSERT INTO imported.business_metadata VALUES('roles',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (json.dumps(sorted(active_roles)),))
        as_of = store.now()
        connection.execute("UPDATE imported.business_metadata SET value=? WHERE key='as_of'", (json.dumps(as_of),))
        if first:
            for kind in ('incidents', 'runs', 'plans', 'actions'):
                store.set_meta(connection, 'archived:' + kind, store.all_items(connection, kind))
                connection.execute(f'DELETE FROM {kind}')
        store.set_meta(connection, 'external_business', True)
        store.set_meta(connection, 'external_roles', sorted(active_roles))
        store.set_meta(connection, 'as_of', as_of)
        dataset['business'].update(status='active', message=f'已按主键接入 {len(records)} 行；业务规则按当前时间分析。')
        store.put(connection, 'datasets', dataset)
        store.audit(connection, 'dataset.activate', f'已接入业务表 {dataset["name"]}。')
    with store.transaction() as connection:
        store.refresh_business(connection)
    return dataset
